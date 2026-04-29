"""Generate user guide Word document for CDSoft Sales Management System."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# ── Styles ────────────────────────────────────────────────────────────────────
styles = doc.styles

def set_heading_style(style_name, font_size, bold=True, color=None, space_before=12, space_after=6):
    try:
        style = styles[style_name]
        style.font.size = Pt(font_size)
        style.font.bold = bold
        if color:
            style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
    except Exception:
        pass

set_heading_style('Heading 1', 16, color=(31, 78, 121), space_before=18, space_after=8)
set_heading_style('Heading 2', 13, color=(21, 101, 192), space_before=14, space_after=6)
set_heading_style('Heading 3', 11, color=(2, 136, 209), space_before=10, space_after=4)

normal = styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11)

# ── Helpers ───────────────────────────────────────────────────────────────────

def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = RGBColor(21, 101, 192)
    return p

def h3(text):
    p = doc.add_heading(text, level=3)
    p.runs[0].font.color.rgb = RGBColor(2, 136, 209)
    return p

def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
    return p

def bullet(text, level=0):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
    return p

def numbered(text, level=0):
    p = doc.add_paragraph(text, style='List Number')
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
    return p

def note(text, note_type='Lưu ý'):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(6)
    border_run = p.add_run(f'⚠ {note_type}: ')
    border_run.bold = True
    border_run.font.size = Pt(11)
    border_run.font.color.rgb = RGBColor(230, 126, 34)
    content_run = p.add_run(text)
    content_run.font.size = Pt(11)
    content_run.font.name = 'Times New Roman'
    return p

def tip(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(6)
    icon_run = p.add_run('✔ Mẹo: ')
    icon_run.bold = True
    icon_run.font.size = Pt(11)
    icon_run.font.color.rgb = RGBColor(39, 174, 96)
    content_run = p.add_run(text)
    content_run.font.size = Pt(11)
    content_run.font.name = 'Times New Roman'
    return p

def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '1F4E79')
        shd.set(qn('w:color'), 'FFFFFF')
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        fill = 'F2F9FF' if row_idx % 2 == 0 else 'FFFFFF'
        for col_idx, cell_text in enumerate(row_data):
            row_cells[col_idx].text = str(cell_text)
            row_cells[col_idx].paragraphs[0].runs[0].font.size = Pt(10)
            row_cells[col_idx].paragraphs[0].runs[0].font.name = 'Times New Roman'
            tc = row_cells[col_idx]._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), fill)
            shd.set(qn('w:val'), 'clear')
            tcPr.append(shd)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    return table

def page_break():
    doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# COVER PAGE
# ══════════════════════════════════════════════════════════════════════════════

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('\n\n\n')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('PHẦN MỀM QUẢN LÝ BÁN HÀNG')
run.font.size = Pt(28)
run.font.bold = True
run.font.color.rgb = RGBColor(31, 78, 121)
run.font.name = 'Times New Roman'

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('CDSoft Sales Management System')
run.font.size = Pt(18)
run.font.bold = False
run.font.color.rgb = RGBColor(21, 101, 192)
run.font.name = 'Times New Roman'

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('─' * 40)
run.font.color.rgb = RGBColor(21, 101, 192)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('HƯỚNG DẪN SỬ DỤNG HỆ THỐNG\nDÀNH CHO NGƯỜI DÙNG CUỐI')
run.font.size = Pt(20)
run.font.bold = True
run.font.color.rgb = RGBColor(31, 78, 121)
run.font.name = 'Times New Roman'

doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(f'Phiên bản: 1.0\nNgày phát hành: {datetime.date.today().strftime("%d/%m/%Y")}')
run.font.size = Pt(12)
run.font.name = 'Times New Roman'
run.font.color.rgb = RGBColor(100, 100, 100)

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# MỤC LỤC (thủ công)
# ══════════════════════════════════════════════════════════════════════════════

h1('MỤC LỤC')

toc_items = [
    ('1.', 'TỔNG QUAN HỆ THỐNG', '3'),
    ('2.', 'ĐĂNG NHẬP HỆ THỐNG', '4'),
    ('3.', 'QUẢN LÝ DANH MỤC — SẢN PHẨM', '5'),
    ('4.', 'QUẢN LÝ KHÁCH HÀNG', '7'),
    ('5.', 'QUẢN LÝ NHÀ CUNG CẤP', '9'),
    ('6.', 'QUẢN LÝ KHO HÀNG', '10'),
    ('  6.1', 'Tồn kho', '10'),
    ('  6.2', 'Phiếu nhập kho', '11'),
    ('  6.3', 'Phiếu xuất kho', '12'),
    ('  6.4', 'Điều chuyển kho', '13'),
    ('  6.5', 'Kiểm kê kho', '14'),
    ('7.', 'QUẢN LÝ ĐƠN HÀNG', '15'),
    ('  7.1', 'Đơn bán hàng', '15'),
    ('  7.2', 'Đơn mua hàng', '17'),
    ('  7.3', 'Khuyến mãi & Voucher', '18'),
    ('  7.4', 'Trả hàng (RMA)', '19'),
    ('8.', 'HOÁ ĐƠN & THANH TOÁN', '20'),
    ('9.', 'CHƯƠNG TRÌNH TÍCH ĐIỂM (LOYALTY)', '22'),
    ('10.', 'BÁO CÁO & THỐNG KÊ', '23'),
    ('11.', 'QUẢN LÝ NGƯỜI DÙNG & PHÂN QUYỀN', '25'),
    ('12.', 'CÀI ĐẶT HỆ THỐNG', '27'),
    ('13.', 'QUẢN LÝ PLATFORM (Dành cho Super Admin)', '28'),
]

for num, title, pg in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    tab_stops = p.paragraph_format.tab_stops
    from docx.shared import Pt as Pt2
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    tab_stops.add_tab_stop(Inches(5.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run_num = p.add_run(f'{num}  {title}')
    run_num.font.size = Pt(11)
    run_num.font.name = 'Times New Roman'
    if not num.startswith(' '):
        run_num.font.bold = True
    run_pg = p.add_run(f'\t{pg}')
    run_pg.font.size = Pt(11)
    run_pg.font.name = 'Times New Roman'

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 1 — TỔNG QUAN
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 1: TỔNG QUAN HỆ THỐNG')

body('CDSoft Sales Management System là phần mềm quản lý bán hàng toàn diện theo mô hình multi-tenant SaaS — cho phép nhiều doanh nghiệp (tenant) vận hành độc lập trên cùng một nền tảng, với dữ liệu hoàn toàn cách ly.')

h2('1.1 Kiến trúc tổng quan')
body('Hệ thống gồm 2 tầng người dùng:')
bullet('Platform Admin (Super Admin): Quản trị viên hệ thống — quản lý các doanh nghiệp sử dụng phần mềm.')
bullet('Tenant User: Người dùng thuộc một doanh nghiệp cụ thể — bao gồm các vai trò: Admin, Giám đốc, Kế toán, Nhân viên kho, Nhân viên bán hàng.')

h2('1.2 Các phân hệ chính')
add_table(
    ['#', 'Phân hệ', 'Mô tả'],
    [
        ['1', 'Danh mục', 'Quản lý sản phẩm, danh mục, đơn vị tính, bảng giá'],
        ['2', 'Khách hàng', 'Hồ sơ khách hàng, hạn mức tín dụng, lịch sử giao dịch'],
        ['3', 'Nhà cung cấp', 'Hồ sơ nhà cung cấp, điều khoản thanh toán'],
        ['4', 'Kho hàng', 'Nhập/xuất/chuyển kho, kiểm kê, tồn kho realtime'],
        ['5', 'Đơn hàng', 'Bán hàng, mua hàng, trả hàng, khuyến mãi'],
        ['6', 'Hoá đơn & Thanh toán', 'Xuất hoá đơn, ghi nhận thanh toán, công nợ'],
        ['7', 'Loyalty', 'Tích điểm, hạng thành viên, đổi điểm'],
        ['8', 'Báo cáo', 'Doanh số, kho, tài chính, KPI nhân viên'],
        ['9', 'Hệ thống', 'Phân quyền, audit log, quản lý người dùng'],
    ],
    [1, 4, 10]
)

h2('1.3 Yêu cầu truy cập')
bullet('Trình duyệt: Chrome, Edge, Firefox phiên bản mới nhất')
bullet('Kết nối internet ổn định')
bullet('Tài khoản được cấp bởi Admin doanh nghiệp')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 2 — ĐĂNG NHẬP
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 2: ĐĂNG NHẬP HỆ THỐNG')

h2('2.1 Đăng nhập Tenant (Người dùng doanh nghiệp)')
body('Truy cập địa chỉ được cung cấp bởi Admin doanh nghiệp.')
numbered('Nhập Email và Mật khẩu.')
numbered('Nhấn nút Đăng nhập.')
numbered('Hệ thống chuyển đến trang Dashboard của doanh nghiệp.')

note('Nếu tài khoản bị khoá hoặc tenant bị tạm ngưng, hệ thống sẽ hiển thị thông báo lỗi. Liên hệ Admin để được hỗ trợ.')

h2('2.2 Đăng xuất')
body('Nhấn nút Logout ở góc trên bên phải màn hình.')

h2('2.3 Chuyển đổi ngôn ngữ')
body('Nhấn vào nút VI / EN ở góc trên phải để chuyển đổi giữa Tiếng Việt và Tiếng Anh. Lựa chọn ngôn ngữ được lưu lại cho lần đăng nhập tiếp theo.')

h2('2.4 Quên mật khẩu')
body('Liên hệ Admin doanh nghiệp để được cấp lại mật khẩu.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 3 — SẢN PHẨM
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 3: QUẢN LÝ DANH MỤC — SẢN PHẨM')

h2('3.1 Danh sách sản phẩm')
body('Vào menu Sản phẩm để xem toàn bộ sản phẩm trong hệ thống.')
bullet('Tìm kiếm theo tên, mã SKU hoặc barcode.')
bullet('Lọc theo Danh mục, Trạng thái (Đang bán / Ngừng bán).')
bullet('Nhấn vào tên sản phẩm để xem chi tiết.')

h2('3.2 Thêm sản phẩm mới')
numbered('Nhấn nút Thêm sản phẩm.')
numbered('Điền thông tin cơ bản:')
bullet('Tên sản phẩm (bắt buộc)', 1)
bullet('Mã SKU (bắt buộc, duy nhất trong hệ thống)', 1)
bullet('Barcode', 1)
bullet('Danh mục', 1)
bullet('Thương hiệu', 1)
bullet('Mô tả', 1)
numbered('Cấu hình đơn vị tính:')
bullet('Đơn vị cơ bản (VD: Cái)', 1)
bullet('Thêm đơn vị chuyển đổi (VD: Thùng = 24 Cái, Hộp = 6 Cái)', 1)
numbered('Cấu hình bảng giá:')
bullet('Giá vốn', 1)
bullet('Giá bán lẻ', 1)
bullet('Giá bán buôn', 1)
bullet('Giá đại lý', 1)
numbered('Cấu hình tồn kho:')
bullet('Kho mặc định', 1)
bullet('Tồn kho tối thiểu (cảnh báo khi hàng sắp hết)', 1)
bullet('Tồn kho tối đa', 1)
numbered('Nhấn Lưu để hoàn tất.')

note('Mã SKU không thể trùng với sản phẩm khác trong cùng doanh nghiệp.')

h2('3.3 Chỉnh sửa sản phẩm')
numbered('Tìm sản phẩm trong danh sách.')
numbered('Nhấn biểu tượng Sửa (bút chì).')
numbered('Cập nhật thông tin cần thay đổi.')
numbered('Nhấn Lưu.')

h2('3.4 Ngừng bán sản phẩm')
body('Nhấn nút Kích hoạt/Ngừng bán trên trang chi tiết sản phẩm. Sản phẩm ngừng bán sẽ không hiển thị trong màn hình tạo đơn hàng nhưng dữ liệu lịch sử vẫn được giữ nguyên.')

h2('3.5 Xoá sản phẩm')
body('Nhấn nút Xoá trên trang chi tiết. Hệ thống chỉ cho phép xoá sản phẩm khi tồn kho = 0 và không có đơn hàng đang xử lý.')

h2('3.6 Quản lý danh mục')
body('Vào menu Danh mục để quản lý cây danh mục sản phẩm.')
bullet('Thêm danh mục cha / danh mục con.')
bullet('Sửa tên danh mục.')
bullet('Xoá danh mục (chỉ khi không có sản phẩm gắn).')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 4 — KHÁCH HÀNG
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 4: QUẢN LÝ KHÁCH HÀNG')

h2('4.1 Danh sách khách hàng')
body('Vào menu Khách hàng để xem danh sách.')
bullet('Tìm kiếm theo tên, mã khách hàng hoặc số điện thoại.')
bullet('Lọc theo nhóm khách hàng (Lẻ / Buôn / Đại lý).')

h2('4.2 Thêm khách hàng mới')
numbered('Nhấn nút Thêm khách hàng.')
numbered('Điền thông tin:')

add_table(
    ['Trường', 'Bắt buộc', 'Mô tả'],
    [
        ['Tên khách hàng', 'Có', 'Tên đầy đủ hoặc tên công ty'],
        ['Mã khách hàng', 'Có', 'Mã duy nhất, tự đặt hoặc auto-generate'],
        ['Số điện thoại', 'Không', 'Dùng để tìm kiếm'],
        ['Email', 'Không', ''],
        ['Địa chỉ', 'Không', 'Địa chỉ giao hàng mặc định'],
        ['Mã số thuế', 'Không', 'Dùng khi xuất hoá đơn VAT'],
        ['Nhóm khách hàng', 'Không', 'Lẻ / Buôn / Đại lý — ảnh hưởng bảng giá'],
        ['Hạn mức tín dụng', 'Không', 'Giới hạn nợ tối đa (0 = không giới hạn)'],
        ['Thời hạn nợ (ngày)', 'Không', 'Số ngày thanh toán sau khi mua nợ'],
        ['Nhân viên phụ trách', 'Không', 'NV bán hàng quản lý khách hàng này'],
    ],
    [4, 2, 9]
)

numbered('Nhấn Lưu.')

h2('4.3 Chi tiết khách hàng')
body('Nhấn vào tên khách hàng để xem trang chi tiết, bao gồm:')
bullet('Thông tin liên lạc và hạn mức tín dụng')
bullet('Lịch sử đơn hàng')
bullet('Lịch sử thanh toán và công nợ hiện tại')
bullet('Điểm tích luỹ và hạng thành viên (nếu bật Loyalty)')

h2('4.4 Hạn mức tín dụng')
body('Khi tạo đơn hàng bán nợ (Credit), hệ thống kiểm tra:')
bullet('Nợ hiện tại + Giá trị đơn mới ≤ Hạn mức tín dụng')
body('Nếu vượt hạn mức, hệ thống hiển thị cảnh báo. Nhân viên có thể vẫn tiếp tục đặt hàng nhưng cần lưu ý.')

note('Hạn mức = 0 nghĩa là không giới hạn nợ. Chỉ áp dụng với đơn thanh toán dạng Công nợ (CREDIT).')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 5 — NHÀ CUNG CẤP
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 5: QUẢN LÝ NHÀ CUNG CẤP')

h2('5.1 Danh sách nhà cung cấp')
body('Vào menu Nhà cung cấp để xem danh sách. Tìm kiếm theo tên hoặc mã nhà cung cấp.')

h2('5.2 Thêm nhà cung cấp mới')
numbered('Nhấn nút Thêm nhà cung cấp.')
numbered('Điền thông tin:')
bullet('Tên nhà cung cấp (bắt buộc)', 1)
bullet('Mã nhà cung cấp', 1)
bullet('Mã số thuế', 1)
bullet('Số điện thoại, Email', 1)
bullet('Địa chỉ', 1)
bullet('Người liên hệ', 1)
bullet('Điều khoản thanh toán (số ngày)', 1)
bullet('Tài khoản ngân hàng', 1)
numbered('Nhấn Lưu.')

h2('5.3 Nhà cung cấp kiêm khách hàng')
body('Nếu một đối tác vừa là nhà cung cấp vừa là khách hàng, bật flag "Cũng là khách hàng" và chọn hồ sơ khách hàng tương ứng. Hệ thống hỗ trợ cấn trừ công nợ 2 chiều.')

h2('5.4 Chi tiết nhà cung cấp')
body('Xem lịch sử nhập hàng, công nợ phải trả và lịch thanh toán dự kiến.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 6 — QUẢN LÝ KHO
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 6: QUẢN LÝ KHO HÀNG')

body('Module quản lý kho bao gồm: Tồn kho, Phiếu nhập, Phiếu xuất, Điều chuyển và Kiểm kê.')
body('Truy cập từ menu Kho hàng → chọn phân hệ tương ứng.')

# 6.1
h2('6.1 Xem tồn kho')
body('Vào menu Kho hàng → Tồn kho để xem số lượng tồn theo từng sản phẩm và kho.')
bullet('Cột Tồn kho hiển thị số lượng thực tế (tính theo đơn vị cơ bản).')
bullet('Cột Đặt trước: số lượng đã được reserve cho đơn hàng CONFIRMED chưa xuất.')
bullet('Hàng nền đỏ: tồn kho ≤ định mức tối thiểu — cần nhập hàng.')
bullet('Lọc theo kho, sản phẩm hoặc nhóm hàng.')

# 6.2
h2('6.2 Phiếu nhập kho')
body('Dùng khi hàng thực tế đã về kho và cần ghi nhận vào hệ thống ngay.')

h3('Tạo phiếu nhập kho mới')
numbered('Vào Kho hàng → Phiếu nhập → Tạo phiếu nhập kho.')
numbered('Chọn Nhà cung cấp và Kho nhập.')
numbered('Thêm sản phẩm:')
bullet('Tìm sản phẩm và nhấn chọn để thêm vào danh sách.', 1)
bullet('Chọn đơn vị tính (VD: Thùng / Hộp / Cái).', 1)
bullet('Nhập số lượng và giá nhập.', 1)
bullet('Nhập thông tin lô hàng (batch number, hạn sử dụng) nếu cần.', 1)
numbered('Nhấn Lưu nháp để lưu tạm, hoặc Xác nhận nhập kho để tồn kho tăng ngay.')

note('Chỉ phiếu ở trạng thái DRAFT mới có thể chỉnh sửa. Sau khi xác nhận, tồn kho đã được cập nhật và không thể huỷ (chỉ điều chỉnh).')

h3('Trạng thái phiếu nhập')
add_table(
    ['Trạng thái', 'Ý nghĩa'],
    [
        ['DRAFT (Nháp)', 'Chưa ảnh hưởng tồn kho, có thể sửa/xoá'],
        ['CONFIRMED (Đã xác nhận)', 'Tồn kho đã tăng, không thể sửa'],
        ['CANCELLED (Đã huỷ)', 'Phiếu bị huỷ, không ảnh hưởng tồn kho'],
    ],
    [4, 11]
)

h3('Chỉnh sửa phiếu nhập (DRAFT)')
body('Mở phiếu nhập → Nhấn Chỉnh sửa → Thay đổi thông tin → Lưu.')

# 6.3
h2('6.3 Phiếu xuất kho')
body('Dùng để ghi nhận hàng xuất ra khỏi kho (nội bộ, bán hàng, điều chuyển nội bộ, v.v.).')

h3('Tạo phiếu xuất kho mới')
numbered('Vào Kho hàng → Phiếu xuất → Tạo phiếu xuất kho.')
numbered('Chọn Kho xuất và Lý do xuất (Bán hàng / Nội bộ / Khác).')
numbered('Thêm sản phẩm cần xuất:')
bullet('Chọn sản phẩm.', 1)
bullet('Chọn đơn vị tính.', 1)
bullet('Nhập số lượng. Hệ thống hiển thị tồn hiện tại để đối chiếu.', 1)
bullet('Nếu số lượng cần xuất vượt tồn thực tế, hệ thống cảnh báo ngay.', 1)
numbered('Nhấn Lưu nháp hoặc Xác nhận xuất kho.')

note('Khi chọn đơn vị chuyển đổi (VD: Thùng), hệ thống tự quy đổi sang đơn vị cơ bản để kiểm tra tồn. VD: Xuất 3 Thùng × 24 Cái = 72 Cái sẽ kiểm tra tồn theo Cái.')

h3('Trạng thái phiếu xuất')
add_table(
    ['Trạng thái', 'Ý nghĩa'],
    [
        ['DRAFT (Nháp)', 'Chưa ảnh hưởng tồn kho, có thể sửa'],
        ['CONFIRMED (Đã xác nhận)', 'Tồn kho đã giảm'],
        ['CANCELLED (Đã huỷ)', 'Nếu huỷ phiếu đã xác nhận, tồn kho được hoàn lại'],
    ],
    [4, 11]
)

# 6.4
h2('6.4 Điều chuyển kho')
body('Dùng khi cần chuyển hàng từ kho này sang kho khác trong cùng doanh nghiệp.')

h3('Tạo lệnh điều chuyển')
numbered('Vào Kho hàng → Điều chuyển → Tạo lệnh điều chuyển.')
numbered('Chọn Kho đi và Kho đến.')
numbered('Thêm sản phẩm và số lượng cần chuyển.')
numbered('Nhấn Tạo lệnh điều chuyển.')

h3('Xác nhận nhận hàng (phía kho đến)')
numbered('Mở lệnh điều chuyển.')
numbered('Nhấn Xác nhận nhận hàng.')
numbered('Kiểm tra số lượng thực nhận, điều chỉnh nếu có chênh lệch.')
numbered('Nhấn Hoàn tất.')

body('Sau khi hoàn tất: tồn kho tại kho đi giảm, tồn kho tại kho đến tăng.')

# 6.5
h2('6.5 Kiểm kê kho')
body('Kiểm kê định kỳ để đối chiếu tồn kho thực tế với sổ sách hệ thống.')

numbered('Vào Kho hàng → Kiểm kê → Tạo phiếu kiểm kê.')
numbered('Chọn kho cần kiểm kê.')
numbered('Hệ thống tự điền số lượng theo sổ sách vào cột Số sách.')
numbered('Nhân viên nhập số lượng thực đếm được vào cột Thực tế.')
numbered('Nhấn Hoàn tất kiểm kê.')
numbered('Hệ thống tự động tạo phiếu điều chỉnh cho các dòng có chênh lệch.')

note('Nên thực hiện kiểm kê ngoài giờ cao điểm để tránh phát sinh giao dịch trong khi đang kiểm đếm.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 7 — ĐƠN HÀNG
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 7: QUẢN LÝ ĐƠN HÀNG')

h2('7.1 Đơn bán hàng')

h3('Tạo đơn bán hàng mới')
numbered('Vào menu Đơn bán hàng → Tạo đơn hàng mới.')
numbered('Chọn khách hàng (tìm theo tên hoặc mã).')
body('Hệ thống tự động hiển thị nhóm khách hàng, nợ hiện tại và hạn mức tín dụng.')
numbered('Chọn kho xuất hàng.')
numbered('Chọn hình thức thanh toán: Tiền mặt / Chuyển khoản / Công nợ.')
numbered('Thêm sản phẩm:')
bullet('Tìm và chọn sản phẩm từ dropdown.', 1)
bullet('Hệ thống tự điền giá theo bảng giá phù hợp nhóm khách hàng.', 1)
bullet('Điều chỉnh số lượng, giá, chiết khấu từng dòng nếu cần.', 1)
numbered('Áp mã voucher (nếu có).')
numbered('Nhấn Lưu nháp hoặc Xác nhận đơn hàng.')

h3('Trạng thái đơn bán hàng')
add_table(
    ['Trạng thái', 'Ý nghĩa', 'Hành động tiếp theo'],
    [
        ['DRAFT', 'Nháp, chưa xác nhận', 'Xác nhận hoặc Huỷ'],
        ['CONFIRMED', 'Đã xác nhận, tồn kho được reserve', 'Giao hàng hoặc Huỷ'],
        ['DELIVERING', 'Đang giao hàng', 'Xác nhận đã giao'],
        ['DELIVERED', 'Đã giao, tồn kho đã trừ', 'Thanh toán / Trả hàng'],
        ['PARTIALLY_RETURNED', 'Trả hàng một phần', '—'],
        ['FULLY_RETURNED', 'Đã trả toàn bộ', '—'],
        ['CANCELLED', 'Đã huỷ', '—'],
    ],
    [3.5, 5.5, 5]
)

h3('Xác nhận đơn hàng')
body('Khi xác nhận, hệ thống kiểm tra tồn kho cho toàn bộ sản phẩm trong đơn. Nếu bất kỳ sản phẩm nào không đủ tồn, hệ thống thông báo chi tiết và không cho phép xác nhận.')

h3('Huỷ đơn hàng')
body('Chỉ huỷ được đơn ở trạng thái DRAFT hoặc CONFIRMED và chưa có thanh toán. Cần nhập lý do huỷ. Tồn kho đã reserve sẽ được hoàn lại.')

h3('Chi tiết đơn hàng')
body('Mở đơn hàng để xem đầy đủ thông tin: danh sách sản phẩm, tổng tiền, trạng thái thanh toán, lịch sử thao tác.')

h2('7.2 Đơn mua hàng')
body('Đơn mua hàng dùng để quản lý quá trình đặt mua từ nhà cung cấp.')

h3('Tạo đơn mua hàng mới')
numbered('Vào menu Đơn mua hàng → Tạo đơn mua hàng.')
numbered('Chọn nhà cung cấp (bắt buộc).')
numbered('Chọn kho nhập hàng.')
numbered('Chọn hình thức thanh toán.')
numbered('Thêm sản phẩm: chọn sản phẩm, nhập số lượng và giá nhập.')
numbered('Nhấn Lưu nháp hoặc Xác nhận đơn hàng.')

h3('Trạng thái đơn mua hàng')
add_table(
    ['Trạng thái', 'Ý nghĩa', 'Hành động tiếp theo'],
    [
        ['DRAFT', 'Nháp', 'Xác nhận đơn'],
        ['CONFIRMED', 'Đã xác nhận với NCC', 'Nhập hàng khi hàng về'],
        ['DELIVERED', 'Đã nhận hàng, tồn kho tăng', 'Thanh toán'],
        ['CANCELLED', 'Đã huỷ', '—'],
    ],
    [3.5, 5.5, 5]
)

note('Tồn kho chỉ tăng sau khi nhấn Nhập hàng (trạng thái chuyển sang DELIVERED), không tăng khi chỉ xác nhận đơn.')

h2('7.3 Khuyến mãi & Voucher')

h3('Chương trình khuyến mãi')
body('Vào menu Cài đặt → Khuyến mãi để quản lý các chương trình:')
bullet('Chiết khấu % hoặc số tiền cố định trên đơn hàng')
bullet('Mua X tặng Y')
bullet('Combo sản phẩm')
bullet('Áp dụng cho nhóm khách hàng cụ thể')
bullet('Đặt ngày bắt đầu / kết thúc chương trình')

h3('Voucher / Mã giảm giá')
body('Tạo voucher tại Cài đặt → Khuyến mãi → tab Voucher:')
bullet('Giảm % hoặc số tiền cố định')
bullet('Đặt điều kiện: giá trị đơn tối thiểu, giới hạn lượt dùng, hạn sử dụng')
body('Khi tạo đơn hàng: nhập mã voucher vào ô Mã khuyến mãi → Áp dụng. Hệ thống tự tính và hiển thị số tiền giảm.')

h2('7.4 Trả hàng (RMA)')
numbered('Mở đơn hàng đã giao (trạng thái DELIVERED).')
numbered('Nhấn nút Trả hàng.')
numbered('Chọn sản phẩm cần trả và nhập số lượng.')
numbered('Nhập lý do trả hàng.')
numbered('Chọn hình thức hoàn tiền: Tiền mặt / Cấn trừ công nợ / Chuyển khoản.')
numbered('Nhấn Xác nhận.')

body('Sau khi trả hàng: tồn kho tại kho được cộng lại, hệ thống tạo phiếu chi hoàn tiền (nếu áp dụng).')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 8 — HOÁ ĐƠN & THANH TOÁN
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 8: HOÁ ĐƠN & THANH TOÁN')

h2('8.1 Hoá đơn')
body('Hoá đơn được tạo tự động khi đơn hàng được xác nhận. Vào menu Thanh toán → Hoá đơn để xem danh sách.')

h3('Xem chi tiết hoá đơn')
bullet('Thông tin doanh nghiệp và khách hàng')
bullet('Danh sách sản phẩm, số lượng, đơn giá, chiết khấu')
bullet('Tổng tiền, số đã thanh toán, số còn nợ')
bullet('Lịch sử các lần thanh toán')

h3('In / Xuất PDF hoá đơn')
body('Nhấn nút In hoặc Export PDF trên trang chi tiết hoá đơn.')

h2('8.2 Ghi nhận thanh toán')
numbered('Mở hoá đơn cần thanh toán.')
numbered('Nhấn nút Ghi thanh toán.')
numbered('Chọn phương thức: Tiền mặt / Chuyển khoản ngân hàng.')
numbered('Nhập số tiền thanh toán (có thể thanh toán một phần).')
numbered('Nhập mã tham chiếu giao dịch (nếu chuyển khoản).')
numbered('Nhấn Xác nhận.')

body('Hệ thống tự động cập nhật số tiền đã trả trên hoá đơn và công nợ khách hàng.')

h2('8.3 Công nợ phải thu (Accounts Receivable)')
body('Vào menu Tài chính → Công nợ phải thu để xem báo cáo tuổi nợ khách hàng.')

add_table(
    ['Nhóm nợ', 'Ý nghĩa'],
    [
        ['0–30 ngày', 'Nợ mới phát sinh, trong hạn thanh toán'],
        ['31–60 ngày', 'Cần đôn đốc'],
        ['61–90 ngày', 'Nợ lâu, rủi ro trung bình'],
        ['>90 ngày', 'Nợ khó đòi, cần xử lý ngay'],
    ],
    [4, 11]
)

h2('8.4 Công nợ phải trả (Accounts Payable)')
body('Vào menu Tài chính → Công nợ phải trả để theo dõi nợ với nhà cung cấp.')
bullet('Xem lịch thanh toán dự kiến.')
bullet('Xem các khoản sắp đến hạn.')
bullet('Cấn trừ công nợ 2 chiều nếu NCC cũng là KH.')

h2('8.5 Quản lý quỹ & Ngân hàng')
body('Vào menu Tài chính → Quỹ & Ngân hàng:')
bullet('Xem số dư quỹ tiền mặt và tài khoản ngân hàng.')
bullet('Tạo phiếu thu/chi thủ công.')
bullet('Xem lịch sử dòng tiền vào/ra.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 9 — LOYALTY
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 9: CHƯƠNG TRÌNH TÍCH ĐIỂM (LOYALTY)')

h2('9.1 Cấu hình chương trình Loyalty')
body('Vào menu Cài đặt → Tích điểm để cấu hình.')

add_table(
    ['Cài đặt', 'Ý nghĩa'],
    [
        ['Tỷ lệ tích điểm', 'VD: Chi 10.000₫ được 1 điểm'],
        ['Giá trị quy đổi', 'VD: 100 điểm = 10.000₫ giảm giá'],
        ['Ngưỡng hạng Bạc', 'Tổng điểm tích luỹ để đạt hạng Bạc'],
        ['Ngưỡng hạng Vàng', 'Tổng điểm tích luỹ để đạt hạng Vàng'],
        ['Ngưỡng hạng Kim cương', 'Tổng điểm tích luỹ để đạt hạng Kim cương'],
        ['Thời hạn điểm', 'Số ngày điểm còn hiệu lực (0 = không hết hạn)'],
        ['Chu kỳ đánh giá hạng', 'Số ngày tính lại hạng thành viên'],
    ],
    [5, 10]
)

h2('9.2 Cách điểm được cộng')
body('Điểm được cộng tự động sau khi khách hàng thanh toán đơn hàng (không cộng cho đơn DRAFT/CONFIRMED). Xem lịch sử điểm trong trang chi tiết khách hàng.')

h2('9.3 Tiêu điểm khi mua hàng')
body('Trong màn hình tạo đơn hàng, nếu khách hàng có điểm tích luỹ, nhân viên có thể áp dụng điểm để giảm giá trực tiếp vào đơn.')

h2('9.4 Hạng thành viên')
add_table(
    ['Hạng', 'Quyền lợi thường dùng'],
    [
        ['Thành viên thường', 'Tích điểm cơ bản'],
        ['Bạc', 'Tích điểm × 1.2, ưu tiên CSKH'],
        ['Vàng', 'Tích điểm × 1.5, giảm giá đặc biệt'],
        ['Kim cương', 'Tích điểm × 2, ưu đãi VIP'],
    ],
    [4, 11]
)

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 10 — BÁO CÁO
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 10: BÁO CÁO & THỐNG KÊ')

h2('10.1 Dashboard tổng quan')
body('Trang chủ sau khi đăng nhập hiển thị các chỉ số quan trọng:')
bullet('Doanh thu hôm nay / tuần này / tháng này')
bullet('Số đơn hàng theo trạng thái')
bullet('Top sản phẩm bán chạy')
bullet('Cảnh báo hàng sắp hết tồn')

h2('10.2 Báo cáo doanh số')
body('Vào menu Báo cáo → Doanh số:')
bullet('Lọc theo khoảng thời gian: ngày / tuần / tháng / quý / năm.')
bullet('Lọc theo nhân viên bán hàng.')
bullet('Xem doanh thu thuần, doanh thu theo sản phẩm, theo nhóm hàng.')
bullet('Biểu đồ xu hướng doanh thu theo thời gian.')
bullet('Bảng top khách hàng theo doanh số.')

h2('10.3 Báo cáo kho (Nhập - Xuất - Tồn)')
body('Vào menu Báo cáo → Kho hàng:')
bullet('Nhập - Xuất - Tồn theo kỳ cho từng sản phẩm.')
bullet('Phân tích ABC: phân loại hàng theo mức độ đóng góp doanh thu.')
bullet('Báo cáo hàng tồn lâu ngày (deadstock).')
bullet('Báo cáo hàng sắp hết hạn sử dụng.')

h2('10.4 Báo cáo tài chính')
body('Vào menu Báo cáo → Tài chính:')

add_table(
    ['Báo cáo', 'Nội dung'],
    [
        ['Lãi lỗ (P&L)', 'Doanh thu thuần, lợi nhuận gộp, chi phí, lợi nhuận ròng'],
        ['Lưu chuyển tiền tệ', 'Dòng tiền từ hoạt động kinh doanh, đầu tư, tài chính'],
        ['Công nợ phải thu', 'Báo cáo tuổi nợ theo khách hàng'],
        ['Công nợ phải trả', 'Lịch thanh toán theo nhà cung cấp'],
    ],
    [4, 11]
)

h2('10.5 Báo cáo KPI nhân viên')
body('Vào menu Báo cáo → KPI:')
bullet('Doanh số thực hiện vs chỉ tiêu của từng nhân viên.')
bullet('Tỷ lệ chốt đơn.')
bullet('Số khách hàng mới.')
bullet('Tỷ lệ nợ quá hạn.')

h2('10.6 Hoa hồng nhân viên')
body('Vào menu Báo cáo → Hoa hồng:')
bullet('Cấu hình công thức tính hoa hồng (% doanh thu hoặc % lợi nhuận).')
bullet('Xem chi tiết hoa hồng từng nhân viên theo kỳ.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 11 — NGƯỜI DÙNG & PHÂN QUYỀN
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 11: QUẢN LÝ NGƯỜI DÙNG & PHÂN QUYỀN')

h2('11.1 Danh sách người dùng')
body('Admin vào menu Người dùng để quản lý tài khoản trong doanh nghiệp.')

h2('11.2 Thêm người dùng mới')
numbered('Nhấn Thêm người dùng.')
numbered('Nhập Họ tên đầy đủ.')
numbered('Nhập Email (dùng để đăng nhập).')
numbered('Đặt Mật khẩu ban đầu.')
numbered('Chọn Vai trò (Role).')
numbered('Nhấn Lưu.')
body('Người dùng có thể đổi mật khẩu sau khi đăng nhập lần đầu.')

h2('11.3 Vai trò mặc định')
add_table(
    ['Vai trò', 'Quyền hạn chính'],
    [
        ['ADMIN', 'Toàn quyền trong doanh nghiệp'],
        ['MANAGER (Giám đốc)', 'Xem tất cả, duyệt phiếu chi, xem báo cáo'],
        ['ACCOUNTANT (Kế toán)', 'Quản lý hoá đơn, thanh toán, công nợ, báo cáo tài chính'],
        ['WAREHOUSE (Nhân viên kho)', 'Nhập/xuất kho, điều chuyển, kiểm kê'],
        ['STAFF (Nhân viên bán hàng)', 'Tạo đơn hàng, xem khách hàng, xem sản phẩm'],
    ],
    [4.5, 10.5]
)

h2('11.4 Tạo vai trò tuỳ chỉnh')
numbered('Vào Cài đặt → Vai trò.')
numbered('Nhấn Thêm vai trò.')
numbered('Đặt tên vai trò.')
numbered('Tích chọn các quyền cụ thể trong bảng phân quyền (theo module và hành động: Xem / Tạo / Sửa / Xoá).')
numbered('Nhấn Lưu.')
numbered('Gán vai trò mới cho người dùng tại màn hình Người dùng.')

h2('11.5 Vô hiệu hoá người dùng')
body('Nhấn nút Vô hiệu hoá trên trang chi tiết người dùng. Người dùng bị vô hiệu hoá không thể đăng nhập nhưng dữ liệu lịch sử vẫn được giữ.')

h2('11.6 Nhật ký hệ thống (Audit Log)')
body('Vào Cài đặt → Nhật ký hệ thống để xem lịch sử tất cả thao tác:')
bullet('Ai đã thực hiện thao tác gì.')
bullet('Thời gian thực hiện.')
bullet('Dữ liệu trước và sau khi thay đổi.')
bullet('Lọc theo người dùng, loại thao tác, khoảng thời gian.')

note('Audit log không thể xoá và chỉ Admin / Manager có quyền xem.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 12 — CÀI ĐẶT HỆ THỐNG
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 12: CÀI ĐẶT HỆ THỐNG')

h2('12.1 Quản lý kho')
body('Vào Cài đặt → Kho hàng để thêm/sửa/xoá kho.')
bullet('Mỗi doanh nghiệp có thể có nhiều kho.')
bullet('Mỗi kho có tên, địa chỉ và trạng thái hoạt động.')

h2('12.2 Cài đặt tài chính')
body('Vào Cài đặt → Tài chính để cấu hình:')
bullet('Quỹ tiền mặt: tên quỹ, số dư ban đầu.')
bullet('Tài khoản ngân hàng: tên ngân hàng, số tài khoản, chủ tài khoản.')

h2('12.3 Theo dõi hàng gần hết hạn')
body('Hệ thống tự động cảnh báo khi hàng sắp hết hạn trong vòng 3–6 tháng. Xem danh sách tại Kho hàng → Cảnh báo hạn sử dụng.')

h2('12.4 Tra cứu bảo hành (Serial Number)')
body('Vào menu Tra cứu bảo hành → Nhập số serial → Xem lịch sử nhập kho, xuất kho và thông tin bảo hành.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 13 — PLATFORM ADMIN
# ══════════════════════════════════════════════════════════════════════════════

h1('CHƯƠNG 13: QUẢN LÝ PLATFORM (Dành cho Super Admin)')

body('Phần này chỉ dành cho quản trị viên hệ thống (Super Admin) — không phải người dùng doanh nghiệp thông thường.')

h2('13.1 Đăng nhập Platform')
body('Truy cập đường dẫn /platform/login và đăng nhập bằng tài khoản Platform Admin.')

h2('13.2 Danh sách doanh nghiệp (Tenant)')
body('Xem và quản lý tất cả doanh nghiệp đang sử dụng hệ thống.')

add_table(
    ['Trạng thái', 'Ý nghĩa'],
    [
        ['PENDING', 'Đang khởi tạo database, chờ hoàn tất'],
        ['ACTIVE', 'Đang hoạt động bình thường'],
        ['SUSPENDED', 'Tạm ngưng, người dùng không thể đăng nhập'],
    ],
    [4, 11]
)

h2('13.3 Thêm doanh nghiệp mới')
numbered('Nhấn Thêm doanh nghiệp.')
numbered('Điền thông tin: Tên công ty, Mã tenant, Tên đăng nhập.')
numbered('Điền email Admin của doanh nghiệp.')
numbered('Nhấn Tạo.')
numbered('Hệ thống tự khởi tạo database riêng và tạo tài khoản Admin.')
numbered('Khi trạng thái chuyển sang ACTIVE, thông báo cho doanh nghiệp để đăng nhập.')

h2('13.4 Tạm ngưng / Kích hoạt doanh nghiệp')
body('Trên trang chi tiết doanh nghiệp, nhấn Tạm ngưng hoặc Kích hoạt. Khi tạm ngưng, toàn bộ người dùng của doanh nghiệp đó không thể đăng nhập.')

h2('13.5 Reset mật khẩu Admin doanh nghiệp')
body('Trên trang chi tiết doanh nghiệp, nhấn Reset mật khẩu Admin. Hệ thống tạo mật khẩu mới và gửi email thông báo.')

h2('13.6 Nhật ký hệ thống Platform')
body('Xem toàn bộ thao tác quản trị: tạo tenant, thay đổi trạng thái, reset mật khẩu, v.v.')

page_break()

# ══════════════════════════════════════════════════════════════════════════════
# PHỤ LỤC
# ══════════════════════════════════════════════════════════════════════════════

h1('PHỤ LỤC: CÂU HỎI THƯỜNG GẶP (FAQ)')

faq = [
    ('Tôi quên mật khẩu, phải làm gì?',
     'Liên hệ Admin doanh nghiệp để được cấp lại mật khẩu. Admin có thể reset mật khẩu trong menu Người dùng.'),
    ('Tại sao không xác nhận được đơn hàng?',
     'Kiểm tra tồn kho: hệ thống không cho phép xác nhận đơn khi tồn kho không đủ. Xem thông báo lỗi để biết sản phẩm nào thiếu và cần nhập thêm bao nhiêu.'),
    ('Tại sao không xoá được sản phẩm?',
     'Sản phẩm chỉ xoá được khi tồn kho = 0 và không có đơn hàng đang xử lý. Hãy điều chỉnh tồn kho hoặc hoàn tất/huỷ các đơn liên quan trước.'),
    ('Điểm tích luỹ của khách hàng không tăng sau khi thanh toán?',
     'Điểm chỉ được cộng sau khi hoá đơn ở trạng thái đã thanh toán đầy đủ. Kiểm tra xem thanh toán đã được ghi nhận chưa trong trang chi tiết hoá đơn.'),
    ('Tại sao không điều chỉnh được phiếu nhập kho?',
     'Chỉ phiếu ở trạng thái DRAFT mới có thể chỉnh sửa. Phiếu đã CONFIRMED đã ảnh hưởng tồn kho. Nếu cần điều chỉnh, hãy tạo phiếu điều chỉnh kho.'),
    ('Cách kiểm tra ai đã thực hiện thao tác nào?',
     'Vào Cài đặt → Nhật ký hệ thống. Lọc theo người dùng và khoảng thời gian cần tra cứu.'),
    ('Làm thế nào để chuyển ngôn ngữ?',
     'Nhấn nút VI / EN ở góc trên bên phải màn hình. Ngôn ngữ thay đổi ngay lập tức và được lưu cho lần sau.'),
]

for i, (q, a) in enumerate(faq, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    q_run = p.add_run(f'Q{i}. {q}')
    q_run.bold = True
    q_run.font.size = Pt(11)
    q_run.font.color.rgb = RGBColor(21, 101, 192)
    q_run.font.name = 'Times New Roman'

    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(0.5)
    p2.paragraph_format.space_after = Pt(4)
    a_run = p2.add_run(f'→ {a}')
    a_run.font.size = Pt(11)
    a_run.font.name = 'Times New Roman'

# Footer
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('─' * 50)
run.font.color.rgb = RGBColor(200, 200, 200)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(f'CDSoft Sales Management System — Hướng dẫn sử dụng v1.0\n© {datetime.date.today().year} CDSoft. All rights reserved.')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(150, 150, 150)
run.font.name = 'Times New Roman'

# ── Save ──────────────────────────────────────────────────────────────────────
output_path = r'd:\Github\cdsoft-sale-management\docs\huong-dan-su-dung.docx'
doc.save(output_path)
print(f'Document saved: {output_path}')
